#!/usr/bin/env python3
"""Build mcp-plugin-scoring.xlsx from shortlist.csv and results.json (see criteria.md for the rubric)."""
import csv
import json
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation

AS_OF = datetime(2026, 10, 15)  # date the WordPress.org data was fetched
FONT = "Arial"

BLUE = Font(name=FONT, size=10, color="0000FF")          # inputs
BLACK = Font(name=FONT, size=10)                           # formulas / text
BOLD = Font(name=FONT, size=10, bold=True)
HEAD = Font(name=FONT, size=10, bold=True, color="FFFFFF")
TITLE = Font(name=FONT, size=14, bold=True)
GREY = Font(name=FONT, size=9, italic=True, color="595959")
INPUT_FILL = PatternFill("solid", fgColor="FFF9DB")        # cells to fill in
HEAD_FILL = PatternFill("solid", fgColor="1F3A5F")
AREA_FILLS = {"setup": "2E75B6", "security": "C0504D", "cap": "548235", "qual": "7F6000", "total": "404040"}
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

SETUP = ["Time to first tool call", "Steps needed / ease", "Works in Claude Desktop (no terminal)", "In-plugin connection guide", "Error messages"]
SECURITY = ["Auth method", "Permission scoping", "Read-only mode", "Destructive-action confirmation", "Audit log", "Safe defaults"]
CAPS = ["Task pass rate (auto)", "Breadth of tools (auto)", "Output quality"]
QUAL = ["Update recency (auto)", "Adoption (auto)", "Rating (auto)", "WP compatibility (auto)", "Abilities API / standards", "Docs & support"]
TASKS = [
    "T1 List 5 most recent posts",
    "T2 Create draft post (heading, paragraph, image)",
    "T3 Edit an existing post's title and content",
    "T4 Upload an image to the media library",
    "T5 List installed plugins and versions",
    "T6 Change a site setting (tagline)",
    "T7 Delete a post: confirmed and logged?",
    "T8 Disallowed action: blocked?",
]
TASK_IDS = [t.split()[0] for t in TASKS]
RESULTS = Path("results.json")  # written by run_tests.py (source: auto) and import_manual.py (source: manual)
AREA_COUNT = 9                   # content areas in criteria.md (Capabilities > breadth), matched by run_tests.py
# Tool names hide what these plugins can do (meta-tools / grouped toolsets), so breadth stays a manual score
BREADTH_UNDERCOUNT = {"cowboy-mcp", "acrossai-mcp-manager"}
AUTH_LABELS = {"basic": "App password", "bearer": "API key", "header": "API key", "none": "None", "oauth": "OAuth"}


def load_results(slugs):
    """Checks per slug from results.json, limited to the shortlist."""
    if not RESULTS.exists():
        return {}
    return {s: c for s, c in json.loads(RESULTS.read_text()).items() if s in slugs}


def check_row(checks):
    """One row of the Checks sheet from a plugin's results.json entries."""
    def get(cid, key="status"):
        return (checks.get(cid) or {}).get(key)
    a3, a3pre, auth = get("A3", "value") or {}, get("A3pre", "value") or {}, get("AUTH", "value") or {}
    if a3.get("hard_fail") or a3pre.get("hard_fail"):
        hard_fail = "Y"
    else:
        hard_fail = "N" if get("A3") in ("Pass", "Public by design", "Fail") else ""
    if get("AUTH") == "N/A":
        method = "External (N/A)"
    else:
        method = AUTH_LABELS.get(auth.get("auth_method"), auth.get("auth_method") or "")
    br = get("BR")
    return {
        "Version": get("S1", "value") or "",
        "MCP routes before enabling": len(get("A2", "value") or []) if get("A2") else "",
        "Out of the box (A3pre)": get("A3pre") or "",
        "Unauthenticated after enabling (A3)": get("A3") or "",
        "Hard fail (auto)": hard_fail,
        "Auth method": method,
        "Tools": auth.get("tool_count") if auth.get("tool_count") is not None else "",
        "Abilities behind meta-tools": auth.get("ability_count") if auth.get("ability_count") is not None else "",
        f"Areas covered (of {AREA_COUNT})": int(br.split("/")[0]) if br else "",
        "Areas": ", ".join(get("BR", "value") or []),
        "Abilities API / MCP Adapter": ", ".join(get("ABIL", "value") or []) or get("ABIL") or "",
        "A3 detail": get("A3", "detail") or "",
        "Last run": max((c.get("at", "") for c in checks.values()), default="")[:10],
    }


def style_header(cell, color=None):
    cell.font = HEAD
    cell.fill = PatternFill("solid", fgColor=color) if color else HEAD_FILL
    cell.alignment = CENTER
    cell.border = BOX


def main():
    plugins = list(csv.DictReader(open("shortlist.csv")))
    scored = [p for p in plugins if p["lane"] != "baseline"]
    baseline = [p for p in plugins if p["lane"] == "baseline"]
    n = len(scored)
    R0 = 5                      # first data row on Scores / Tasks
    R1 = R0 + n - 1             # last data row
    P0 = 5                      # first data row on Plugins
    P1 = P0 + len(plugins) - 1
    results = load_results({p["slug"] for p in scored})

    wb = Workbook()

    # ---------------------------------------------------------------- README
    ws = wb.active
    ws.title = "README"
    ws["A1"] = "WordCamp Joburg: MCP plugin comparison scoring sheet"
    ws["A1"].font = TITLE
    lines = [
        ("How to use", True),
        ("1. For each plugin, install it on a clean site (setup_test_sites.sh), connect Claude Desktop, then work through the Tasks sheet (Pass / Partial / Fail).", False),
        ("2. On the Scores sheet, fill the yellow cells with a score from 0 to 5 (blue text = your input). Grey-header columns marked (auto) are formulas, so leave them alone. 'Breadth of tools (auto)' is yellow where there is no automated result: score those by hand.", False),
        ("3. 'Security hard fail? (auto)' comes from the unauthenticated probe on the Checks sheet: Y when the enabled endpoint returns tools without credentials and the source does not show it is public by design. The plugin's total is then capped (Weights sheet).", False),
        ("This workbook is regenerated from shortlist.csv and results.json (python3 build_scoring_sheet.py), so record results in results.json rather than typing them here.", False),
        ("4. Rank appears once both Setup and Security have at least one score. Weights, thresholds and the cap live on the Weights sheet.", False),
        ("", False),
        ("Scoring scale (0-5)", True),
        ("0 = missing or broken   1 = poor   2 = weak   3 = acceptable   4 = good   5 = excellent", False),
        ("Example: 'In-plugin connection guide' = 5 if the plugin shows a ready-to-copy Claude Desktop config; 3 if there is generic text only; 0 if nothing.", False),
        ("", False),
        ("Colour key", True),
        ("Blue text on yellow = input.   Black = formula or fixed data.   Weights and thresholds on the Weights sheet are inputs too.", False),
        ("", False),
        ("Sheets", True),
        ("Weights: area weights (must total 100), hard-fail cap, thresholds for the auto-scored columns.", False),
        ("Plugins: WordPress.org data for each plugin (fetched " + AS_OF.strftime("%Y-%m-%d") + "), plus the mcp-adapter baseline reference row.", False),
        ("Scores: one row per plugin; weighted total out of 100 and rank.", False),
        ("Tasks: the 8 standard tasks, run via Claude Desktop, per plugin. Results already in results.json are filled in.", False),
        ("Checks: automated results from run_tests.py (results.json): out-of-the-box state, unauthenticated probe, auth method, tool count, breadth, Abilities API use.", False),
        ("", False),
        ("Baseline", True),
        ("mcp-adapter is the official WordPress MCP Adapter. It is listed on the Plugins sheet as the reference point for comparison and is not scored or ranked.", False),
        ("", False),
        ("Lanes", True),
        ("established = 100+ installs.  challenger = under 100 installs, hand-picked after a pre-screen, so newer plugins get a fair chance. Both lanes are scored identically.", False),
    ]
    for i, (t, b) in enumerate(lines, start=3):
        c = ws.cell(row=i, column=1, value=t)
        c.font = BOLD if b else BLACK
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 130

    # --------------------------------------------------------------- Weights
    wt = wb.create_sheet("Weights")
    wt["A1"] = "Weights and thresholds"
    wt["A1"].font = TITLE
    wt["A3"], wt["B3"] = "Data as of", AS_OF
    wt["B3"].number_format = "yyyy-mm-dd"
    wt["C3"] = "Date the WordPress.org data was fetched; used for 'days since update'."
    wt["A5"], wt["B5"] = "Area", "Weight"
    for c in ("A5", "B5"):
        style_header(wt[c])
    areas = [("Setup & connection", 30), ("Security & auth", 30), ("Capabilities", 25), ("Quality & maintenance", 15)]
    for i, (a, w) in enumerate(areas, start=6):
        wt.cell(row=i, column=1, value=a)
        wt.cell(row=i, column=2, value=w)
    wt["A10"], wt["B10"] = "Total", "=SUM(B6:B9)"
    wt["C10"] = '=IF(B10=100,"OK","Weights must total 100")'
    wt["A12"] = "Security hard-fail cap"
    wt["A13"], wt["B13"] = "Max total score if 'Security hard fail?' = Y", 40
    wt["C13"] = "Assumption: set by the presenter (starting point, adjust as you like)."
    wt["A15"] = "Update recency thresholds (days since last update, for 5 / 4 / 3 / 2 points; older = 1)"
    for i, v in enumerate((30, 90, 180, 365), start=16):
        wt.cell(row=i, column=1, value=f"{5 - (i - 16)} points if updated within (days)")
        wt.cell(row=i, column=2, value=v)
    wt["A21"] = "Adoption thresholds (active installs, for 5 / 4 / 3 / 2 points; fewer = 1)"
    for i, v in enumerate((10000, 1000, 100, 10), start=22):
        wt.cell(row=i, column=1, value=f"{5 - (i - 22)} points if installs at least")
        wt.cell(row=i, column=2, value=v)
        wt.cell(row=i, column=2).number_format = "#,##0"
    wt["A27"] = "Note: adoption is one of six Quality sub-scores (15% of the total), so low installs cost a challenger little."
    wt["A27"].font = GREY
    for r in range(3, 26):
        for col in (1, 2, 3):
            c = wt.cell(row=r, column=col)
            if c.font == Font() or c.font.name != FONT:
                c.font = BLACK
    for ref in ["B3"] + [f"B{r}" for r in (6, 7, 8, 9, 13, 16, 17, 18, 19, 22, 23, 24, 25)]:
        wt[ref].font = BLUE
        wt[ref].fill = INPUT_FILL
        wt[ref].border = BOX
    for ref in ("A12", "A15", "A21", "A10", "B10"):
        wt[ref].font = BOLD
    wt["C3"].font = GREY
    wt["C13"].font = GREY
    wt.column_dimensions["A"].width = 78
    wt.column_dimensions["B"].width = 14
    wt.column_dimensions["C"].width = 70

    # --------------------------------------------------------------- Plugins
    pl = wb.create_sheet("Plugins")
    pl["A1"] = "Shortlisted plugins (WordPress.org data)"
    pl["A1"].font = TITLE
    pl["A2"] = "Data fetched " + AS_OF.strftime("%Y-%m-%d") + ". Rating is 0-100 as returned by the API. The baseline row is a reference and is not scored."
    pl["A2"].font = GREY
    heads = ["Lane", "Slug", "Name", "Active installs", "Last updated", "Days since update", "Rating (0-100)", "Num ratings",
             "Tested up to current WP?", "Mentions Abilities API / MCP Adapter?", "URL", "Short description"]
    for j, h in enumerate(heads, start=1):
        style_header(pl.cell(row=4, column=j, value=h))
    for i, p in enumerate(plugins, start=P0):
        vals = [p["lane"], p["slug"], p["name"], int(p["active_installs"] or 0),
                datetime.strptime(p["last_updated"], "%Y-%m-%d"), f"=Weights!$B$3-E{i}",
                int(p["rating"] or 0), int(p["num_ratings"] or 0), p["tested_current"], p["mentions_abilities"], p["url"], p["short_description"]]
        for j, v in enumerate(vals, start=1):
            c = pl.cell(row=i, column=j, value=v)
            c.font = BLACK
            c.border = BOX
            c.alignment = Alignment(vertical="top", wrap_text=(j in (3, 12)))
        pl.cell(row=i, column=4).number_format = "#,##0"
        pl.cell(row=i, column=5).number_format = "yyyy-mm-dd"
        pl.cell(row=i, column=6).number_format = "0"
        pl.cell(row=i, column=11).hyperlink = p["url"]
        pl.cell(row=i, column=11).font = Font(name=FONT, size=10, color="0563C1", underline="single")
    for j, w in enumerate([12, 36, 40, 12, 12, 10, 10, 10, 12, 16, 44, 80], start=1):
        pl.column_dimensions[L(j)].width = w
    pl.row_dimensions[4].height = 42
    pl.freeze_panes = "D5"

    # ----------------------------------------------------------------- Tasks
    tk = wb.create_sheet("Tasks")
    tk["A1"] = "Standard task results (run via Claude Desktop)"
    tk["A1"].font = TITLE
    tk["A2"] = "Enter Pass, Partial or Fail. Partial counts as half. T7 and T8 test safety behaviour: Pass = confirmed/logged (T7) or correctly blocked (T8)."
    tk["A2"].font = GREY
    th = ["Slug", "Name"] + TASKS + ["Pass rate", "Tasks run"]
    for j, h in enumerate(th, start=1):
        style_header(tk.cell(row=4, column=j, value=h))
    tk.row_dimensions[4].height = 66
    dv_task = DataValidation(type="list", formula1='"Pass,Partial,Fail"', allow_blank=True, showErrorMessage=True,
                             errorTitle="Invalid result", error="Use Pass, Partial or Fail.")
    tk.add_data_validation(dv_task)
    c0, c1 = 3, 2 + len(TASKS)           # task columns C..J
    pass_col, run_col = c1 + 1, c1 + 2   # K, L
    for i, p in enumerate(scored, start=R0):
        tk.cell(row=i, column=1, value=p["slug"]).font = BLACK
        tk.cell(row=i, column=2, value=p["name"]).font = BLACK
        tk.cell(row=i, column=2).alignment = Alignment(wrap_text=True, vertical="top")
        rng = f"{L(c0)}{i}:{L(c1)}{i}"
        for j, tid in zip(range(c0, c1 + 1), TASK_IDS):
            status = (results.get(p["slug"], {}).get(tid) or {}).get("status")
            c = tk.cell(row=i, column=j, value=status if status in ("Pass", "Partial", "Fail") else None)
            c.font, c.fill, c.border, c.alignment = BLUE, INPUT_FILL, BOX, Alignment(horizontal="center")
            dv_task.add(c)
        tk.cell(row=i, column=pass_col, value=f'=IF(COUNTA({rng})=0,"",(COUNTIF({rng},"Pass")+0.5*COUNTIF({rng},"Partial"))/COUNTA({rng}))').number_format = "0%"
        tk.cell(row=i, column=run_col, value=f"=COUNTA({rng})")
        for j in (1, 2, pass_col, run_col):
            tk.cell(row=i, column=j).border = BOX
            if j in (pass_col, run_col):
                tk.cell(row=i, column=j).font = BLACK
                tk.cell(row=i, column=j).alignment = Alignment(horizontal="center")
    task_rng = f"{L(c0)}{R0}:{L(c1)}{R1}"
    tk.conditional_formatting.add(task_rng, CellIsRule(operator="equal", formula=['"Pass"'], fill=PatternFill("solid", bgColor="C6EFCE")))
    tk.conditional_formatting.add(task_rng, CellIsRule(operator="equal", formula=['"Partial"'], fill=PatternFill("solid", bgColor="FFEB9C")))
    tk.conditional_formatting.add(task_rng, CellIsRule(operator="equal", formula=['"Fail"'], fill=PatternFill("solid", bgColor="FFC7CE")))
    tk.column_dimensions["A"].width = 36
    tk.column_dimensions["B"].width = 36
    for j in range(c0, c1 + 1):
        tk.column_dimensions[L(j)].width = 15
    tk.column_dimensions[L(pass_col)].width = 10
    tk.column_dimensions[L(run_col)].width = 10
    tk.freeze_panes = "C5"

    # ---------------------------------------------------------------- Checks
    ck = wb.create_sheet("Checks")
    ck["A1"] = "Automated checks (run_tests.py)"
    ck["A1"].font = TITLE
    ck["A2"] = ("From results.json. Out of the box = before MCP is switched on; A3 = after switching it on and enabling every tool. "
                "Breadth matches tool and ability names, so it undercounts generic tools. Blank = not run or needs the manual sheet (OAuth, external server).")
    ck["A2"].font = GREY
    rows = {p["slug"]: check_row(results.get(p["slug"], {})) for p in scored}
    ck_heads = ["Slug", "Name"] + list(check_row({}))
    for j, h in enumerate(ck_heads, start=1):
        style_header(ck.cell(row=4, column=j, value=h))
    ck_col = {h: j for j, h in enumerate(ck_heads, start=1)}
    for i, p in enumerate(scored, start=R0):
        for j, v in enumerate([p["slug"], p["name"]] + list(rows[p["slug"]].values()), start=1):
            c = ck.cell(row=i, column=j, value=v)
            c.font, c.border = BLACK, BOX
            c.alignment = Alignment(vertical="top", wrap_text=ck_heads[j - 1] in ("Name", "Areas", "Abilities API / MCP Adapter", "A3 detail"),
                                    horizontal=None if j <= 2 or ck_heads[j - 1] in ("Areas", "A3 detail") else "center")
    hf_rng = f"{L(ck_col['Hard fail (auto)'])}{R0}:{L(ck_col['Hard fail (auto)'])}{R1}"
    ck.conditional_formatting.add(hf_rng, CellIsRule(operator="equal", formula=['"Y"'], fill=PatternFill("solid", bgColor="FFC7CE")))
    for h, j in ck_col.items():
        ck.column_dimensions[L(j)].width = {"Slug": 34, "Name": 36, "Areas": 40, "Abilities API / MCP Adapter": 24, "A3 detail": 70}.get(h, 13)
    ck.row_dimensions[4].height = 54
    ck.freeze_panes = "C5"

    def check_lookup(head, i):
        cl = L(ck_col[head])
        return f"INDEX(Checks!${cl}${R0}:${cl}${R1},MATCH($B{i},Checks!$A${R0}:$A${R1},0))"

    # ---------------------------------------------------------------- Scores
    sc = wb.create_sheet("Scores")
    sc["A1"] = "Weighted scores (out of 100)"
    sc["A1"].font = TITLE
    sc["A2"] = "Yellow cells: score 0-5. Area score = average of that area's sub-scores / 5 x area weight. Blank sub-scores are ignored."
    sc["A2"].font = GREY

    col = 4
    groups = []  # (key, title, first col, last col, score col)
    def block(key, title, names):
        nonlocal col
        first = col
        col += len(names)
        groups.append((key, title, first, col - 1, col, names))
        col += 1  # area score column
    block("setup", "Setup & connection", SETUP)
    block("security", "Security & auth", SECURITY)
    block("cap", "Capabilities", CAPS)
    block("qual", "Quality & maintenance", QUAL)
    hf, raw, fin, rank, notes = col, col + 1, col + 2, col + 3, col + 4

    for j, h in ((1, "Lane"), (2, "Slug"), (3, "Name")):
        style_header(sc.cell(row=4, column=j, value=h))
    for key, title, first, last, scol, names in groups:
        color = AREA_FILLS[key]
        sc.cell(row=3, column=first, value=title)
        for j in range(first, scol + 1):
            style_header(sc.cell(row=3, column=j), color)
        for j, nm in zip(range(first, last + 1), names):
            style_header(sc.cell(row=4, column=j, value=nm), "7F7F7F" if "(auto)" in nm else color)
        style_header(sc.cell(row=4, column=scol, value="Area score"), color)
    for j, h in ((hf, "Security hard fail? (auto)"), (raw, "Raw total"), (fin, "Final total"), (rank, "Rank"), (notes, "Notes")):
        style_header(sc.cell(row=4, column=j, value=h), AREA_FILLS["total"])
    br_col = groups[2][2] + CAPS.index("Breadth of tools (auto)")
    sc.row_dimensions[4].height = 66

    dv_score = DataValidation(type="whole", operator="between", formula1="0", formula2="5", allow_blank=True,
                              showErrorMessage=True, errorTitle="Invalid score", error="Enter a whole number from 0 to 5.")
    sc.add_data_validation(dv_score)

    def lookup(colletter, i):
        return f"INDEX(Plugins!${colletter}${P0}:${colletter}${P1},MATCH($B{i},Plugins!$B${P0}:$B${P1},0))"

    area_weight_row = {"setup": 6, "security": 7, "cap": 8, "qual": 9}
    for i, p in enumerate(scored, start=R0):
        for j, v in ((1, p["lane"]), (2, p["slug"]), (3, p["name"])):
            c = sc.cell(row=i, column=j, value=v)
            c.font, c.border = BLACK, BOX
            c.alignment = Alignment(vertical="top", wrap_text=(j == 3))
        breadth_auto = rows[p["slug"]][f"Areas covered (of {AREA_COUNT})"] != "" and p["slug"] not in BREADTH_UNDERCOUNT
        for key, title, first, last, scol, names in groups:
            for j, nm in zip(range(first, last + 1), names):
                c = sc.cell(row=i, column=j)
                c.border = BOX
                c.alignment = Alignment(horizontal="center")
                if "(auto)" in nm and (j != br_col or breadth_auto):
                    c.font = BLACK
                else:
                    c.font, c.fill = BLUE, INPUT_FILL
                    dv_score.add(c)
            rng = f"{L(first)}{i}:{L(last)}{i}"
            a = sc.cell(row=i, column=scol, value=f"=IF(COUNT({rng})=0,0,AVERAGE({rng})/5*Weights!$B${area_weight_row[key]})")
            a.number_format, a.font, a.border = "0.0", BOLD, BOX
        # auto formulas
        capcol = groups[2][2]
        sc.cell(row=i, column=capcol, value=f'=IFERROR(IF(INDEX(Tasks!${L(pass_col)}${R0}:${L(pass_col)}${R1},MATCH($B{i},Tasks!$A${R0}:$A${R1},0))="","",INDEX(Tasks!${L(pass_col)}${R0}:${L(pass_col)}${R1},MATCH($B{i},Tasks!$A${R0}:$A${R1},0))*5),"")').number_format = "0.0"
        if breadth_auto:
            sc.cell(row=i, column=br_col, value=f"={check_lookup(f'Areas covered (of {AREA_COUNT})', i)}/{AREA_COUNT}*5").number_format = "0.0"
        q0 = groups[3][2]
        days = lookup("F", i)
        sc.cell(row=i, column=q0, value=f"=IF({days}<=Weights!$B$16,5,IF({days}<=Weights!$B$17,4,IF({days}<=Weights!$B$18,3,IF({days}<=Weights!$B$19,2,1))))")
        inst = lookup("D", i)
        sc.cell(row=i, column=q0 + 1, value=f"=IF({inst}>=Weights!$B$22,5,IF({inst}>=Weights!$B$23,4,IF({inst}>=Weights!$B$24,3,IF({inst}>=Weights!$B$25,2,1))))")
        sc.cell(row=i, column=q0 + 2, value=f'=IF({lookup("H", i)}=0,"",{lookup("G", i)}/20)').number_format = "0.0"
        sc.cell(row=i, column=q0 + 3, value=f'=IF({lookup("I", i)}="Y",5,2)')
        setup_s, sec_s, cap_s, qual_s = (g[4] for g in groups)
        sc.cell(row=i, column=hf, value=f'=IFERROR(""&{check_lookup("Hard fail (auto)", i)},"")').font = BLACK
        sc.cell(row=i, column=raw, value=f"={L(setup_s)}{i}+{L(sec_s)}{i}+{L(cap_s)}{i}+{L(qual_s)}{i}")
        sc.cell(row=i, column=fin, value=f'=IF({L(hf)}{i}="Y",MIN({L(raw)}{i},Weights!$B$13),{L(raw)}{i})')
        ready = f"AND(COUNT({L(groups[0][2])}{i}:{L(groups[0][3])}{i})>0,COUNT({L(groups[1][2])}{i}:{L(groups[1][3])}{i})>0)"
        sc.cell(row=i, column=rank, value=f'=IF({ready},RANK({L(fin)}{i},${L(fin)}${R0}:${L(fin)}${R1}),"")')
        for j in (raw, fin):
            sc.cell(row=i, column=j).number_format = "0.0"
            sc.cell(row=i, column=j).font = BOLD
        sc.cell(row=i, column=rank).font = BOLD
        for j in (hf, raw, fin, rank, notes):
            sc.cell(row=i, column=j).border = BOX
            sc.cell(row=i, column=j).alignment = Alignment(horizontal="center", vertical="top") if j != notes else WRAP
        sc.cell(row=i, column=notes).font, sc.cell(row=i, column=notes).fill = BLUE, INPUT_FILL
    sc.conditional_formatting.add(f"{L(hf)}{R0}:{L(hf)}{R1}", CellIsRule(operator="equal", formula=['"Y"'], fill=PatternFill("solid", bgColor="FFC7CE")))
    sc.column_dimensions["A"].width = 12
    sc.column_dimensions["B"].width = 34
    sc.column_dimensions["C"].width = 36
    for j in range(4, notes):
        sc.column_dimensions[L(j)].width = 13
    sc.column_dimensions[L(notes)].width = 50
    sc.freeze_panes = "D5"

    # order sheets
    wb._sheets = [wb["README"], wb["Scores"], wb["Tasks"], wb["Checks"], wb["Plugins"], wb["Weights"]]
    wb.save("mcp-plugin-scoring.xlsx")
    print(f"saved: {n} plugins on Scores/Tasks (rows {R0}-{R1}), {len(baseline)} baseline on Plugins, "
          f"{len(results)} with results.json checks")


if __name__ == "__main__":
    main()
