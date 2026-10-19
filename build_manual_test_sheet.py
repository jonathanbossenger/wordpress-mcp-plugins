#!/usr/bin/env python3
"""Build mcp-manual-test-sheet.xlsx: copy/paste runbook for the checks that are hard to automate.

Pick a plugin slug on the Runbook sheet and every command updates to that plugin's Studio site.
Record outcomes on the Results sheet (same Pass / Partial / Fail scale as the scoring sheet).
"""
import csv

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation

FONT = "Arial"
BLACK = Font(name=FONT, size=10)
BOLD = Font(name=FONT, size=10, bold=True)
BLUE = Font(name=FONT, size=10, color="0000FF")
HEAD = Font(name=FONT, size=10, bold=True, color="FFFFFF")
TITLE = Font(name=FONT, size=14, bold=True)
GREY = Font(name=FONT, size=9, italic=True, color="595959")
CODE = Font(name="Consolas", size=10)
INPUT_FILL = PatternFill("solid", fgColor="FFF9DB")
HEAD_FILL = PatternFill("solid", fgColor="1F3A5F")
SECTION_FILL = PatternFill("solid", fgColor="D9E2F3")
CODE_FILL = PatternFill("solid", fgColor="F2F2F2")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

# Input cells on the Runbook sheet (column B)
INPUTS = {
    "slug": "$B$4",
    "root": "$B$5",
    "path": "$B$6",
    "url": "$B$7",
    "endpoint": "$B$8",
    "user": "$B$9",
    "apppass": "$B$10",
    "b64": "$B$11",
    "postid": "$B$12",
    "lpuser": "$B$13",
    "lppass": "$B$14",
}

WP = 'studio wp --path "{path}"'

# (id, area, title, what to do, command template or None, what to record)
SECTIONS = [
    ("A. Before you connect anything (Setup & Security)", [
        ("A1", "Setup", "Start the site",
         "Starts the Studio site for this plugin. Note the URL it prints and paste it into B7. Start your stopwatch for 'time to first tool call' now.",
         "./setup_test_sites.sh start {slug}",
         "Time started."),
        ("A2", "Security", "Probe for an open endpoint (before any config)",
         "Do this BEFORE touching the plugin's settings, so you see the safe-defaults state. First list REST routes that look like MCP, then paste the endpoint into B8.",
         WP + " eval 'foreach (array_keys(rest_get_server()->get_routes()) as $r) {{ if (preg_match(\"/mcp/i\", $r)) echo $r . PHP_EOL; }}'",
         "MCP route(s) found. If none, note where the plugin says the endpoint lives (settings page, readme)."),
        ("A3", "Security", "Unauthenticated tools/list",
         "Call the endpoint with NO credentials. A JSON list of tools means the plugin exposes everything by default: set 'Security hard fail?' to Y on the Scores sheet. A 401/403 is the good outcome. If the endpoint is missing or says MCP is disabled, pick 'Disabled' (good safe default) and run this step again after enabling MCP in A4, because the real hard-fail check is whether the ENABLED endpoint answers without credentials.",
         "curl -si -X POST '{endpoint}' -H 'Content-Type: application/json' -H 'Accept: application/json, text/event-stream' -d '{{\"jsonrpc\":\"2.0\",\"id\":1,\"method\":\"tools/list\"}}' | head -40",
         "HTTP status and whether any tools were returned. Feeds: Safe defaults, Auth method."),
        ("A4", "Setup", "Find the settings and connection UI",
         "Open the plugin's admin page. Does it show a ready-to-copy Claude Desktop config or connector URL? Take a screenshot of the connection instructions and of every security-related setting (read-only, per-tool toggles, roles, confirmations, logging).",
         "U='{url}/wp-admin/admin.php?page={slug}'; open \"$U\" 2>/dev/null || xdg-open \"$U\"",
         "Feeds: In-plugin connection guide, Permission scoping, Read-only mode, Audit log. If the guessed page URL is wrong, use the plugin's menu entry."),
        ("A5", "Security", "Dump the plugin's options",
         "Shows stored settings, which often reveals defaults (enabled flags, API keys, tool allow-lists) the UI hides.",
         WP + " option list --search='*{slug}*' --format=table; " + WP + " option list --search='*mcp*' --format=table",
         "Anything notable: keys stored in plain text, everything enabled by default, etc."),
    ]),
    ("B. Credentials and Claude Desktop connection (Setup)", [
        ("B1", "Setup", "Create an application password",
         "Many plugins accept a WordPress application password. If the plugin uses its own API key or OAuth instead, follow its settings page and note that in Results. Paste the password it prints into B10.",
         WP + " user application-password create {user} mcp-test --porcelain",
         "Which auth method the plugin actually uses (feeds: Auth method)."),
        ("B2", "Setup", "Base64 for the Authorization header",
         "Run the command, paste the output into B11. Only needed for plugins that accept Basic auth.",
         "printf '%s' '{user}:{apppass}' | base64 -w0",
         "n/a"),
        ("B3", "Setup", "Claude Desktop config (stdio via mcp-remote)",
         "Paste into claude_desktop_config.json (Settings > Developer > Edit Config), merge into mcpServers, then fully quit and reopen Claude Desktop. Studio sites are http://localhost, which mcp-remote accepts. If the plugin supports remote connectors/OAuth directly, try Settings > Connectors > Add custom connector first and note whether it worked without a terminal.",
         '{{"mcpServers":{{"{slug}":{{"command":"npx","args":["-y","mcp-remote","{endpoint}","--header","Authorization:${{AUTH_HEADER}}"],"env":{{"AUTH_HEADER":"Basic {b64}"}}}}}}}}',
         "Stop the stopwatch at the first successful tool call. Record minutes, steps needed, and any error messages (feeds: Time to first tool call, Steps, Works without terminal, Error messages)."),
        ("B4", "Setup", "Deliberately break the connection",
         "Change one character of the password (or stop the site) and retry a prompt in Claude Desktop. Is the error understandable?",
         None,
         "Quote the error text. Feeds: Error messages."),
    ]),
    ("C. Standard tasks: paste the prompt into Claude Desktop, then verify with the command (Capabilities)", [
        ("T1", "Task", "List the 5 most recent posts",
         "Prompt: \"List my 5 most recent posts with their titles and status.\"",
         WP + " post list --posts_per_page=5 --orderby=date --order=DESC --fields=ID,post_title,post_status",
         "Does Claude's answer match the command output? Pass / Partial / Fail."),
        ("T2", "Task", "Create a draft post with heading, paragraph and image",
         "Prompt: \"Create a DRAFT post titled 'MCP Test Draft' with a heading, a short paragraph and this image: https://s.w.org/style/images/about/WordPress-logotype-wmark.png\"",
         WP + " post list --post_status=draft --s='MCP Test Draft' --fields=ID,post_title,post_status; " + WP + " post get $(" + WP + " post list --post_status=draft --s='MCP Test Draft' --field=ID | head -1) --field=post_content",
         "Is it a draft (not published)? Is the content valid block markup (<!-- wp:heading -->...) or raw HTML? Is the image present? Open the post in the editor: any 'Attempt recovery' / invalid-block warnings?"),
        ("T3", "Task", "Edit an existing post's title and content",
         "Pick an ID from the first command and paste it into B12. Prompt: \"Change the title of post {postid} to 'Edited by MCP' and add a closing sentence to the content.\"",
         WP + " post list --post_status=publish --posts_per_page=3 --fields=ID,post_title; " + WP + " post get {postid} --fields=ID,post_title,post_modified",
         "Title changed? Original content preserved? Pass / Partial / Fail."),
        ("T4", "Task", "Upload an image to the media library",
         "Prompt: \"Upload this image to the media library: https://s.w.org/style/images/about/WordPress-logotype-wmark.png\" (some plugins can't fetch URLs; note that and try attaching a file if Claude Desktop allows).",
         WP + " post list --post_type=attachment --fields=ID,post_title,post_mime_type",
         "Attachment exists with the right mime type? Pass / Partial / Fail."),
        ("T5", "Task", "List installed plugins and versions",
         "Prompt: \"List the installed plugins with their versions and whether they're active.\"",
         WP + " plugin list --fields=name,version,status",
         "Does the answer match? Some plugins can't do this at all: that is a Fail, not a Partial."),
        ("T6", "Task", "Change a site setting (tagline)",
         "Prompt: \"Change the site tagline to 'Changed by MCP test'.\"",
         WP + " option get blogdescription",
         "Pass if the value changed. Reset afterwards with the cleanup step in section E."),
        ("T7", "Task", "Delete a post: confirmed and logged?",
         "Prompt: \"Delete post {postid}.\" Watch what happens: did Claude/the plugin ask for confirmation, or just do it? Was it trashed or permanently deleted? Then check for any audit trail.",
         WP + " post list --post_status=trash,any --posts_per_page=5 --fields=ID,post_title,post_status; " + WP + " db query \"SHOW TABLES LIKE '%log%'\"; " + WP + " option list --search='*log*' --format=table",
         "Pass = asked for confirmation AND an audit record exists (plugin log page or table). Partial = one of the two. Note trash vs permanent delete."),
        ("T8", "Task", "Disallowed action: blocked?",
         "Create a low-privilege user, create an application password for them (paste user name into B13 and password into B14), reconnect Claude Desktop using THOSE credentials (repeat B2/B3 with the new values), then use the prompt. Prompt: \"Publish a new post titled 'Subscriber Attack' and delete post {postid}.\"",
         WP + " user create {lpuser} {lpuser}@example.test --role=subscriber --user_pass=Sub-pass-12345; " + WP + " user application-password create {lpuser} mcp-lowpriv --porcelain; " + WP + " post list --post_status=any --s='Subscriber Attack' --fields=ID,post_title,post_status",
         "Pass = both actions refused and the post list shows no 'Subscriber Attack'. Fail = the subscriber could write or delete (also consider a Security hard fail)."),
    ]),
    ("D. Quality checks (Capabilities, Quality & maintenance)", [
        ("D1", "Quality", "Does it use the Abilities API / MCP Adapter?",
         "Searches the plugin source for the WordPress Abilities API or the MCP Adapter. Feeds: Abilities API / standards.",
         "grep -rIl --include=*.php -E 'wp_register_ability|WP\\\\MCP|mcp-adapter' \"{path}/wp-content/plugins/{slug}\" | head -10",
         "Files found (yes) or nothing (no). Also check composer.json / vendor/ for wordpress/mcp-adapter."),
        ("D2", "Quality", "Count and skim the tool list",
         "In Claude Desktop, open the plugin's tools list (the tools icon in the chat box). Count tools and note which areas are covered: posts, pages, media, taxonomies, users, plugins, settings, blocks, WooCommerce. Feeds: Breadth of tools.",
         None,
         "Tool count and areas covered."),
        ("D3", "Quality", "Docs and support",
         "Open the support forum and readme. Are questions answered? Is there setup documentation beyond the readme? Feeds: Docs & support.",
         "for U in 'https://wordpress.org/support/plugin/{slug}/' 'https://wordpress.org/plugins/{slug}/#description'; do open \"$U\" 2>/dev/null || xdg-open \"$U\"; done",
         "Last reply date from the maintainer, number of unresolved threads, quality of docs."),
        ("D4", "Quality", "Console / debug log errors",
         "Catches PHP notices and fatals triggered while you ran the tasks.",
         "tail -50 \"{path}/wp-content/debug.log\" 2>/dev/null || echo 'no debug.log'",
         "Any PHP errors attributable to the plugin."),
    ]),
    ("E. Cleanup between plugins", [
        ("E1", "Cleanup", "Reset test data and tagline",
         "Removes content created by the tasks and restores the tagline.",
         WP + " post delete $(" + WP + " post list --post_type=post,attachment --post_status=any --s='MCP Test' --format=ids) --force; " + WP + " option update blogdescription 'Just another MCP test site'",
         "n/a"),
        ("E2", "Cleanup", "Stop the site",
         "Frees the port and memory. Remove the plugin's entry from claude_desktop_config.json and restart Claude Desktop before the next plugin.",
         "./setup_test_sites.sh stop {slug}",
         "n/a"),
    ]),
]

RESULT_COLS = [
    "A3 Unauth probe (Pass = 401/403 once enabled)", "A3 Hard fail?", "Auth method", "Minutes to first tool call", "Steps needed", "Works w/o terminal?",
    "T1", "T2", "T3", "T4", "T5", "T6", "T7", "T8",
    "T7 confirmed?", "T7 audit log?", "Abilities API / MCP Adapter?", "Tool count", "Output quality (0-5)", "Docs & support (0-5)",
]


def formula(template):
    """Turn a command template with {name} placeholders into an Excel concatenation formula."""
    parts, buf, i = [], "", 0
    while i < len(template):
        if template.startswith("{{", i):
            buf += "{"
            i += 2
        elif template.startswith("}}", i):
            buf += "}"
            i += 2
        elif template[i] == "{":
            j = template.index("}", i)
            name = template[i + 1:j]
            if buf:
                parts.append('"' + buf.replace('"', '""') + '"')
                buf = ""
            parts.append(INPUTS[name])
            i = j + 1
        else:
            buf += template[i]
            i += 1
    if buf:
        parts.append('"' + buf.replace('"', '""') + '"')
    # Excel allows ~8k chars in a formula but only 255 per string literal: split long literals
    out = []
    for p in parts:
        if p.startswith('"') and len(p) > 250:
            body = p[1:-1]
            chunk = []
            k = 0
            while k < len(body):
                piece = body[k:k + 200]
                # don't split an escaped quote pair
                if piece.endswith('"') and (len(piece) - len(piece.rstrip('"'))) % 2:
                    piece = body[k:k + 201]
                chunk.append('"' + piece + '"')
                k += len(piece)
            out.extend(chunk)
        else:
            out.append(p)
    return "=" + "&".join(out)


def head(cell, fill=None):
    cell.font, cell.fill = HEAD, fill or HEAD_FILL
    cell.alignment, cell.border = CENTER, BOX


def main():
    plugins = [p for p in csv.DictReader(open("shortlist.csv")) if p["lane"] != "baseline"]
    slugs = [p["slug"] for p in plugins]

    wb = Workbook()

    # ---------------------------------------------------------------- README
    ws = wb.active
    ws.title = "README"
    ws["A1"] = "WordCamp Joburg: MCP plugin manual test sheet"
    ws["A1"].font = TITLE
    lines = [
        ("What this is", True),
        ("The steps that are hard to automate (connecting Claude Desktop, judging settings screens, running prompts, checking safety behaviour), as copy/paste commands.", False),
        ("Everything automatable (site creation, plugin install, WordPress.org metadata scoring) stays in setup_test_sites.sh and the scoring sheet.", False),
        ("", False),
        ("How to use", True),
        ("1. On Runbook, pick the plugin slug in B4 and fill the yellow inputs as you learn them (site URL, endpoint, passwords). Every grey command cell updates automatically.", False),
        ("2. Work down the Runbook in order. Click a grey command cell, copy it (Ctrl+C), paste into a terminal, run it from the repo folder. Prompts for Claude Desktop are in the instructions column.", False),
        ("3. Record the outcome for that plugin on the Results sheet. T1-T8 use Pass / Partial / Fail; copy them into the Tasks sheet of mcp-plugin-scoring.xlsx.", False),
        ("4. Run section E before moving to the next plugin.", False),
        ("", False),
        ("Assumptions to check", True),
        ("Studio sites live in $HOME/Studio/mcp-tests/<slug> ($HOME is used so it expands inside double quotes) (MCP_TEST_ROOT in setup_test_sites.sh). Change B5 if yours differ.", False),
        ("Commands use the 'studio wp' wrapper that setup_test_sites.sh uses, and assume you run them from this repo folder (for ./setup_test_sites.sh).", False),
        ("Command cells are formulas. Copying one from Excel or Google Sheets copies its text. If a cell shows a formula instead, copy from the formula bar.", False),
        ("The Claude Desktop config uses mcp-remote (npx, so Node.js is required) with the Authorization header passed via env, which works around the space-escaping bug on Windows. It only suits plugins that accept Basic auth. Plugins with OAuth or their own API key need the config from their settings page.", False),
        ("The admin page URL in A4 is a guess based on the slug. Use the plugin's own menu entry when it is wrong.", False),
        ("Prompts mention the same image URL for T2/T4; plugins that cannot fetch remote URLs score lower on T4, not a Fail by default. Use your judgement and note it.", False),
        ("", False),
        ("Colour key", True),
        ("Blue text on yellow = input. Grey monospace = command to copy. Pass / Partial / Fail are colour-coded on Results.", False),
    ]
    for i, (t, b) in enumerate(lines, start=3):
        c = ws.cell(row=i, column=1, value=t)
        c.font = BOLD if b else BLACK
        c.alignment = WRAP
    ws.column_dimensions["A"].width = 140

    # --------------------------------------------------------------- Runbook
    rb = wb.create_sheet("Runbook")
    rb["A1"] = "Runbook: one plugin at a time"
    rb["A1"].font = TITLE
    rb["A2"] = "Pick a slug, fill the yellow cells as you go, then copy the grey commands."
    rb["A2"].font = GREY
    labels = [
        (4, "Plugin slug", slugs[0], "Dropdown. Everything below follows this."),
        (5, "Sites root folder", "$HOME/Studio/mcp-tests", "Matches MCP_TEST_ROOT in setup_test_sites.sh."),
        (6, "Site path", f"={INPUTS['root']}&\"/\"&{INPUTS['slug']}", "Formula. Quoted when used in commands."),
        (7, "Site URL", "http://localhost:8881", "Printed by step A1. No trailing slash."),
        (8, "MCP endpoint URL", "", "From step A2 or the plugin's settings page, e.g. http://localhost:8881/wp-json/<namespace>/mcp"),
        (9, "Admin user", "admin", "Created by setup_test_sites.sh."),
        (10, "Application password", "", "Output of step B1 (local throwaway site only)."),
        (11, "Base64 of user:password", "", "Output of step B2."),
        (12, "Test post ID", "", "Pick from the T3 command output."),
        (13, "Low-privilege user", "mcp-subscriber", "Used by T8."),
        (14, "Low-privilege app password", "", "Output of the T8 command."),
    ]
    for r, label, val, note in labels:
        rb.cell(row=r, column=1, value=label).font = BOLD
        c = rb.cell(row=r, column=2, value=val)
        c.border = BOX
        if r == 6:
            c.font = BLACK
        else:
            c.font, c.fill = BLUE, INPUT_FILL
        n = rb.cell(row=r, column=3, value=note)
        n.font = GREY
    dv_slug = DataValidation(type="list", formula1=f"=Slugs!$A$2:$A${len(slugs) + 1}", allow_blank=False)
    rb.add_data_validation(dv_slug)
    dv_slug.add("B4")

    row = 16
    for col, h in enumerate(["ID", "Area", "Step", "What to do", "Copy / paste", "What to record"], start=1):
        head(rb.cell(row=row, column=col, value=h))
    row += 1
    for section, steps in SECTIONS:
        for col in range(1, 7):
            rb.cell(row=row, column=col).fill = SECTION_FILL
        rb.cell(row=row, column=1, value=section).font = BOLD
        row += 1
        for sid, area, title, todo, cmd, record in steps:
            vals = [sid, area, title, todo, formula(cmd) if cmd else "(no command: done in the UI)", record]
            for col, v in enumerate(vals, start=1):
                c = rb.cell(row=row, column=col, value=v)
                c.border, c.alignment = BOX, WRAP
                c.font = BOLD if col == 1 else BLACK
            cc = rb.cell(row=row, column=5)
            if cmd:
                cc.font, cc.fill = CODE, CODE_FILL
            else:
                cc.font = GREY
            row += 1
    for col, w in zip("ABCDEF", (14, 11, 34, 62, 80, 52)):
        rb.column_dimensions[col].width = w
    rb.column_dimensions["B"].width = 30  # inputs share column B with the Area column
    rb.freeze_panes = "A17"

    # --------------------------------------------------------------- Results
    rs = wb.create_sheet("Results")
    rs["A1"] = "Manual results: one row per plugin"
    rs["A1"].font = TITLE
    rs["A2"] = ("T1-T8: Pass / Partial / Fail, same as the Tasks sheet in the scoring workbook. "
                "Scores for 'Output quality' and 'Docs & support' are 0-5 and map to the Scores sheet. Use Notes for anything you want to say on stage.")
    rs["A2"].font = GREY
    cols = ["Slug"] + RESULT_COLS + ["Notes"]
    for j, h in enumerate(cols, start=1):
        head(rs.cell(row=4, column=j, value=h))
    rs.row_dimensions[4].height = 60
    dv_pf = DataValidation(type="list", formula1='"Pass,Partial,Fail"', allow_blank=True)
    dv_yn = DataValidation(type="list", formula1='"Y,N"', allow_blank=True)
    dv_auth = DataValidation(type="list", formula1='"App password,API key,OAuth,None,Other"', allow_blank=True)
    dv_score = DataValidation(type="whole", operator="between", formula1="0", formula2="5", allow_blank=True)
    dv_probe = DataValidation(type="list", formula1='"Pass,Disabled,Fail"', allow_blank=True)
    for dv in (dv_pf, dv_yn, dv_auth, dv_score, dv_probe):
        rs.add_data_validation(dv)
    kinds = {
        "A3 Unauth probe (Pass = 401/403 once enabled)": dv_probe, "A3 Hard fail?": dv_yn, "Auth method": dv_auth, "Works w/o terminal?": dv_yn,
        "T7 confirmed?": dv_yn, "T7 audit log?": dv_yn, "Abilities API / MCP Adapter?": dv_yn,
        "Output quality (0-5)": dv_score, "Docs & support (0-5)": dv_score,
    }
    for t in ("T1", "T2", "T3", "T4", "T5", "T6", "T7", "T8"):
        kinds[t] = dv_pf
    for i, slug in enumerate(slugs, start=5):
        c = rs.cell(row=i, column=1, value=slug)
        c.font, c.border = BLACK, BOX
        for j, name in enumerate(RESULT_COLS + ["Notes"], start=2):
            c = rs.cell(row=i, column=j)
            c.font, c.fill, c.border = BLUE, INPUT_FILL, BOX
            c.alignment = WRAP if name == "Notes" else Alignment(horizontal="center", vertical="top")
            if name in kinds:
                kinds[name].add(c)
    last = 4 + len(slugs)
    task_rng = f"{L(2 + RESULT_COLS.index('T1'))}5:{L(2 + RESULT_COLS.index('T8'))}{last}"
    rs.conditional_formatting.add(task_rng, CellIsRule(operator="equal", formula=['"Pass"'], fill=PatternFill("solid", bgColor="C6EFCE")))
    rs.conditional_formatting.add(task_rng, CellIsRule(operator="equal", formula=['"Partial"'], fill=PatternFill("solid", bgColor="FFEB9C")))
    rs.conditional_formatting.add(task_rng, CellIsRule(operator="equal", formula=['"Fail"'], fill=PatternFill("solid", bgColor="FFC7CE")))
    rs.column_dimensions["A"].width = 38
    for j in range(2, len(cols)):
        rs.column_dimensions[L(j)].width = 14
    rs.column_dimensions[L(len(cols))].width = 60
    rs.freeze_panes = "B5"

    # ----------------------------------------------------------------- Slugs
    sl = wb.create_sheet("Slugs")
    sl["A1"] = "slug"
    sl["A1"].font = BOLD
    for i, s in enumerate(slugs, start=2):
        sl.cell(row=i, column=1, value=s).font = BLACK
    sl.column_dimensions["A"].width = 44

    wb._sheets = [wb["README"], wb["Runbook"], wb["Results"], wb["Slugs"]]
    wb.save("mcp-manual-test-sheet.xlsx")
    print(f"saved mcp-manual-test-sheet.xlsx: {len(slugs)} plugins, {sum(len(s) for _, s in SECTIONS)} steps")


if __name__ == "__main__":
    main()
